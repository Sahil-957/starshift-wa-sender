import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Contacts"

headers = ["Sr No", "Name", "Mobile Number", "Country Code", "Group Name", "Custom1", "Custom2"]
ws.append(headers)

for col, h in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col)
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color="075E54", end_color="075E54", fill_type="solid")
    cell.alignment = Alignment(horizontal="center")

sample_rows = [
    [1, "Rahul Sharma", "8873520027", "91", "", "Pune", "Gold Member"],
    [2, "Priya Patil", "9822011223", "91", "", "Mumbai", "Silver Member"],
    [3, "Sales Team Group", "", "91", "WA Sender Updates", "", ""],
    [4, "Amit Verma", "9004455667", "91", "", "Nagpur", "New Customer"],
]
for row in sample_rows:
    ws.append(row)

widths = [8, 20, 16, 12, 20, 14, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

notes = wb.create_sheet("Instructions")
lines = [
    "WA Bulk Sender - Excel Template Instructions",
    "",
    "Column meanings:",
    "Sr No        -> Serial number (optional, for your reference only)",
    "Name         -> Contact name. Used in message as {{name}}",
    "Mobile Number-> WITHOUT country code, e.g. 8873520027 (leave blank if this row is a Group)",
    "Country Code -> e.g. 91 for India (combined with Mobile Number to form 918873520027)",
    "Group Name   -> Fill this ONLY if you want to send to a WhatsApp Group by its name (leave Mobile Number blank)",
    "Custom1      -> Any extra field for personalization, used as {{custom1}} in message",
    "Custom2      -> Any extra field for personalization, used as {{custom2}} in message",
    "",
    "Message template example:",
    "  Hi {{name}}, greetings from {{custom2}}! Special offer for you in {{custom1}}.",
    "",
    "Rules:",
    "- Each row is either an individual contact (Mobile Number filled) OR a Group (Group Name filled), not both.",
    "- Do not change header names in row 1.",
    "- Save file as .xlsx before uploading to the extension dashboard.",
]
for i, line in enumerate(lines, start=1):
    notes.cell(row=i, column=1, value=line)
notes.column_dimensions["A"].width = 100

wb.save(r"c:\Users\Admin\Desktop\Whatsapp_bulksms_sender_extention\templates\bulk_contacts_template.xlsx")
print("saved")
